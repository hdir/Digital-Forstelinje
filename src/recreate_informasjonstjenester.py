from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "informasjonstjenester-redigerbar.pptx"

PX_TO_IN = 1 / 100

WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(32, 29, 38)
BLUE = RGBColor(0, 88, 180)
BLUE_DARK = RGBColor(0, 83, 174)
BLUE_LIGHT = RGBColor(199, 222, 250)
BLUE_PALE = RGBColor(207, 226, 249)
GREY = RGBColor(218, 218, 218)
SHADOW = RGBColor(150, 150, 150)


def inches(value):
    return Inches(value * PX_TO_IN)


def add_shape(slide, x, y, width, height, fill, radius=True, line=None, shadow=False):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    if shadow:
        shadow_shape = slide.shapes.add_shape(
            shape_type, inches(x + 3), inches(y + 4), inches(width), inches(height)
        )
        shadow_shape.fill.solid()
        shadow_shape.fill.fore_color.rgb = SHADOW
        shadow_shape.line.fill.background()
    shape = slide.shapes.add_shape(shape_type, inches(x), inches(y), inches(width), inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.18
    return shape


def add_text(slide, text, x, y, width, height, size=16, color=BLACK, bold=False,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, font="Arial"):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(width), inches(height))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = 0.94
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def tile(slide, text, x, y, width, height, fill=BLUE_PALE, color=BLACK,
         size=16, bold=True, overlay=None, overlay_width=0):
    add_shape(slide, x, y, width, height, fill, shadow=True)
    if overlay:
        add_shape(slide, x, y, overlay_width, height, overlay, radius=False)
    add_text(slide, text, x + 5, y + 4, width - 10, height - 8, size, color, bold)


def section(slide, title, x, y, width, height):
    add_shape(slide, x, y, width, height, WHITE, radius=False)
    add_text(slide, title, x + 8, y + 1, width - 16, 28, 15, BLUE, False, PP_ALIGN.LEFT)


def build_presentation():
    prs = Presentation()
    prs.slide_width = inches(1471)
    prs.slide_height = inches(972)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE

    # Main rounded frame and title.
    add_shape(slide, 0, 8, 1222, 955, BLUE_LIGHT, shadow=False)
    add_text(slide, "Informasjonstjenester", 265, 28, 670, 45, 28, BLUE_DARK, True)

    # Top information area.
    section(slide, "Skaffe seg oversikt over innbyggeres tilstand og behov for helsehjelp", 28, 75, 1168, 239)
    tile(slide, "Klinisk\noppsum-\nering", 49, 105, 170, 90, size=15)
    tile(slide, "Problem/\ndiagnose og\nbehov", 235, 105, 170, 90, size=15)
    tile(slide, "Plan", 421, 105, 170, 90, fill=BLUE, color=WHITE, size=16)
    tile(slide, "Pågående og\ngjennomførte\nprosedyrer og\nbehandlinger", 607, 105, 170, 90, size=14)
    tile(slide, "Tjenester,\nytelser og\nhjelpemidler", 789, 105, 172, 90, size=14)
    tile(slide, "Legemidler og\nvaksiner", 53, 206, 169, 90, size=15, color=WHITE, overlay=BLUE, overlay_width=87)
    tile(slide, "Immunisering\n(status)", 235, 206, 170, 90, size=14)
    tile(slide, "Kritisk\ninformasjon", 421, 206, 170, 90, fill=BLUE, color=WHITE, size=15)

    section(slide, "Gjøre oppslag i tidligere journalopplysninger", 28, 324, 1168, 136)
    tile(slide, "Undersøkelser,\nmålinger og\nfunn", 49, 360, 170, 90, size=14)
    tile(slide, "Multimedia og\nMTU-målinger", 235, 360, 170, 90, size=15)
    tile(slide, "Journal-\ndokumenter", 421, 360, 170, 90, size=15)
    tile(slide, "Kliniske\nbakgrunns-\nopplysninger", 607, 360, 170, 90, size=14)

    section(slide, "Anmode om eller bestille tjenester eller ytelser, med svar samt kommunisere om saker", 28, 469, 1168, 136)
    tile(slide, "Bestilling og\nsvar (lab)", 49, 505, 173, 90, fill=GREY, size=14)
    tile(slide, "Henvisning\nepikrise, m.m.", 235, 505, 173, 90, fill=GREY, size=14)
    tile(slide, "Anmodning\nom tjeneste", 421, 505, 173, 90, fill=GREY, size=14)
    tile(slide, "Kommunika-\nsjon ved saks-\nbehandling", 607, 505, 173, 90, fill=GREY, size=13)

    # Lower four sections.
    add_shape(slide, 0, 628, 1222, 4, BLUE, radius=False)
    section(slide, "Innhente innbyggeres opplysninger", 28, 646, 578, 136)
    tile(slide, "Pasient-\ndemografi", 49, 680, 170, 90, size=15)
    tile(slide, "Personvern", 235, 680, 170, 90, fill=BLUE, color=WHITE, size=15)
    tile(slide, "Innbyggeres\nopplysninger\nog ønsker", 421, 680, 170, 90, fill=BLUE, color=WHITE, size=13)

    section(slide, "Rapportere egen aktivitet", 622, 646, 574, 136)
    tile(slide, "Rapportering\nhelsefag", 640, 682, 168, 86, fill=GREY, size=14)
    tile(slide, "Rapportering\nadministrativt", 826, 682, 168, 86, fill=GREY, size=14)

    section(slide, "Slå opp i generelle informasjonskilder (grunndata)", 28, 790, 578, 136)
    tile(slide, "Oversikt over\ntilgjengelige\ntjenester og\ntilbud", 49, 824, 170, 90, fill=BLUE, color=WHITE, size=13)
    tile(slide, "Klinisk\nkunnskap", 235, 824, 170, 90, size=15)

    section(slide, "Arrangere og delta i møter, konsultasjoner og samtaler", 622, 790, 574, 136)
    tile(slide, "Team- og møte-\nadministrasjon", 640, 830, 172, 88, fill=GREY, size=13)
    tile(slide, "Video", 826, 830, 172, 88, fill=WHITE, size=15)
    tile(slide, "Tekstlig\ndialog", 1007, 830, 172, 88, fill=GREY, size=14)

    # Right-side communication actions.
    tile(slide, "Sende og\nmotta", 1296, 303, 168, 90, fill=GREY, size=14)
    tile(slide, "Slå opp og\ntilgjengelig-\ngjøre", 1296, 404, 168, 90, size=14)
    tile(slide, "Endre og\ndele", 1296, 505, 168, 90, fill=BLUE, color=WHITE, size=15)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_presentation()